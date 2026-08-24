from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any


RESULT_HEADERS = [
    "Case ID",
    "Status",
    "Error",
    "Has reference",
    "M1 mIoU",
    "M1 mIoU no background",
    "M1 Pixel accuracy",
    "M2 mIoU",
    "M2 mIoU no background",
    "M2 Pixel accuracy",
    "DAPS mIoU",
    "DAPS mIoU no background",
    "DAPS Pixel accuracy",
    "Agreement M1-M2 mIoU",
    "Agreement M1-M2 mIoU no background",
    "Agreement M1-M2 Pixel accuracy",
    "Agreement DAPS-M1 mIoU",
    "Agreement DAPS-M1 mIoU no background",
    "Agreement DAPS-M1 Pixel accuracy",
    "Agreement DAPS-M2 mIoU",
    "Agreement DAPS-M2 mIoU no background",
    "Agreement DAPS-M2 Pixel accuracy",
]


def _metric_value(
    result: dict[str, Any],
    source: str,
    metric: str,
) -> float | None:
    """Estrae una metrica annidata dal risultato del batch."""

    metrics = result.get(source)

    if metrics is None:
        return None

    return metrics.get(metric)


def write_excel_report(
    output_path: str | Path,
    results: list[dict[str, Any]],
    run_info: dict[str, Any],
) -> Path:
    """Crea un report Excel con risultati e metadati della run."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    # Foglio 1: risultati per ciascun caso.
    results_sheet = workbook.active
    results_sheet.title = "Results"
    results_sheet.append(RESULT_HEADERS)

    for result in results:
        results_sheet.append(
            [
                result["case_id"],
                result["status"],
                result["error"],
                result["has_reference"],

                _metric_value(result, "model_1", "miou"),
                _metric_value(result, "model_1", "miou_no_bg"),
                _metric_value(result, "model_1", "pixel_accuracy"),

                _metric_value(result, "model_2", "miou"),
                _metric_value(result, "model_2", "miou_no_bg"),
                _metric_value(result, "model_2", "pixel_accuracy"),

                _metric_value(result, "daps", "miou"),
                _metric_value(result, "daps", "miou_no_bg"),
                _metric_value(result, "daps", "pixel_accuracy"),

                _metric_value(result, "agreement_m1_m2", "miou"),
                _metric_value(result, "agreement_m1_m2", "miou_no_bg"),
                _metric_value(result, "agreement_m1_m2", "pixel_accuracy"),

                _metric_value(result, "agreement_daps_m1", "miou"),
                _metric_value(result, "agreement_daps_m1", "miou_no_bg"),
                _metric_value(result, "agreement_daps_m1", "pixel_accuracy"),

                _metric_value(result, "agreement_daps_m2", "miou"),
                _metric_value(result, "agreement_daps_m2", "miou_no_bg"),
                _metric_value(result, "agreement_daps_m2", "pixel_accuracy"),
            ]
        )

    header_fill = PatternFill("solid", fgColor="0F766E")

    for cell in results_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    results_sheet.freeze_panes = "A2"
    results_sheet.auto_filter.ref = results_sheet.dimensions
    results_sheet.sheet_view.showGridLines = False

    # Tutte le colonne delle metriche: da E a V.
    for row in results_sheet.iter_rows(
        min_row=2,
        min_col=5,
        max_col=22,
    ):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "0.0000"

    results_sheet.column_dimensions["A"].width = 24
    results_sheet.column_dimensions["B"].width = 14
    results_sheet.column_dimensions["C"].width = 45
    results_sheet.column_dimensions["D"].width = 15

    for column in "EFGHIJKLMNOPQRSTUV":
        results_sheet.column_dimensions[column].width = 20

    # Foglio 2: informazioni della run.
    info_sheet = workbook.create_sheet("Run info")
    info_sheet.append(["Property", "Value"])

    for key, value in run_info.items():
        info_sheet.append([key, value])

    info_sheet.append(
        ["Created at", datetime.now().isoformat(timespec="seconds")]
    )

    for cell in info_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")

    info_sheet.column_dimensions["A"].width = 28
    info_sheet.column_dimensions["B"].width = 80
    info_sheet.sheet_view.showGridLines = False

    workbook.save(output_path)
    return output_path